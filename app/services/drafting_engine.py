"""Legal notice and pleading generator.

Templates fix the structure and the standard legal wording of each document type, with
jurisdiction-specific variants (an Indian legal notice, a US demand letter, an English
letter before claim...). Two kinds of fields:

  {{variable}}      a fact (names, addresses, FIR number, deadline). Filled, in order of
                    precedence, from values the lawyer typed, then from the case context by
                    Claude (only when the context states them), then automatic values (date,
                    governing law). Anything still unknown renders as a highlighted
                    "[●description●]" placeholder and is listed as missing information.
  {{ai:slot}}       a passage Claude writes from the case context (facts, grounds, demands).

Case context comes from the case thread (its memory and recent session summaries/facts),
attached documents and anything the lawyer adds. It is PII-masked before Claude sees it and
the result is unmasked. Drafts export to .docx (python-docx) and .pdf (ReportLab), with
placeholders highlighted so nothing unfinished goes out unnoticed.
"""

import io
import re
from dataclasses import dataclass, field
from datetime import date
from xml.sax.saxutils import escape

from app.services.jurisdiction_service import JurisdictionProfile, US_STATES

_FIELD = re.compile(r"\{\{\s*(ai:)?([a-z_]+)\s*\}\}")
PLACEHOLDER = re.compile(r"\[●[^●\]]*●\]")


@dataclass(frozen=True)
class Template:
    doc_type: str
    family: str  # IN | US | UK | * (any jurisdiction)
    name: str
    title: str
    variables: dict[str, str]  # name -> description
    slots: dict[str, str]  # name -> what to write
    sections: list[tuple[str | None, list[str]]]
    signature: list[str] = field(default_factory=list)


DOC_TYPES = {
    "legal_notice": "Legal notice / demand letter",
    "bail_application": "Bail application",
    "cease_and_desist": "Cease and desist letter",
    "nda": "Non-disclosure agreement",
    "affidavit": "Affidavit",
}

FACTS = "Numbered paragraphs (1., 2., ...) setting out the relevant facts in chronological order."

TEMPLATES: list[Template] = [
    # ------------------------------------------------------------------ legal notices
    Template(
        "legal_notice", "IN", "Legal notice (India)", "LEGAL NOTICE",
        {
            "advocate_name": "Advocate's name", "advocate_address": "Advocate's office address",
            "client_name": "Client's full name", "client_description": "Client's description, e.g. 'son of ..., residing at ...'",
            "recipient_name": "Recipient's name", "recipient_address": "Recipient's address",
            "subject": "Short subject of the notice", "compliance_days": "Days allowed to comply (e.g. 15)",
        },
        {
            "facts": FACTS,
            "legal_grounds": "Numbered paragraphs on the recipient's legal liability, citing the applicable statutes.",
            "demands": "What the recipient must do (pay, return, rectify, cease), stated precisely.",
        },
        [
            (None, [
                "From: {{advocate_name}}, Advocate, {{advocate_address}}",
                "Date: {{date}}",
                "To: {{recipient_name}}, {{recipient_address}}",
                "By Registered Post A.D. / Speed Post / Email",
                "Subject: Legal notice on behalf of {{client_name}} regarding {{subject}}",
            ]),
            (None, [
                "Sir/Madam,",
                "Under instructions from and on behalf of my client, {{client_name}}, {{client_description}}, "
                "I hereby serve upon you the following legal notice:",
            ]),
            ("Facts", ["{{ai:facts}}"]),
            ("Legal position", ["{{ai:legal_grounds}}"]),
            ("Demand", [
                "{{ai:demands}}",
                "You are hereby called upon to comply with the above within {{compliance_days}} days of receipt of this "
                "notice, failing which my client shall be constrained to initiate appropriate civil and/or criminal "
                "proceedings against you before the competent court or forum, entirely at your risk as to costs and consequences.",
            ]),
            (None, ["A copy of this notice has been retained in my office for record and further action."]),
        ],
        ["{{advocate_name}}", "Advocate for {{client_name}}"],
    ),
    Template(
        "legal_notice", "US", "Demand letter (US)", "DEMAND LETTER",
        {
            "attorney_name": "Attorney's name", "firm_name": "Law firm", "client_name": "Client's name",
            "recipient_name": "Recipient's name", "recipient_address": "Recipient's address",
            "subject": "Matter / subject", "response_deadline": "Response deadline (a date)",
        },
        {"facts": FACTS, "legal_grounds": "The legal basis of the claim, citing applicable law.", "demands": "The specific demand."},
        [
            (None, ["{{date}}", "{{recipient_name}}", "{{recipient_address}}", "Re: {{subject}}", "Dear {{recipient_name}}:",
                    "This firm represents {{client_name}} in connection with {{subject}}."]),
            ("Background", ["{{ai:facts}}"]),
            ("Legal basis", ["{{ai:legal_grounds}}"]),
            ("Demand", [
                "{{ai:demands}}",
                "If we do not receive a satisfactory response by {{response_deadline}}, our client is prepared to pursue all "
                "available remedies without further notice. Nothing in this letter waives any of our client's rights or "
                "remedies, all of which are expressly reserved.",
            ]),
        ],
        ["Sincerely,", "{{attorney_name}}", "{{firm_name}}"],
    ),
    Template(
        "legal_notice", "UK", "Letter before claim (England & Wales)", "LETTER BEFORE CLAIM",
        {
            "solicitor_name": "Solicitor's name", "firm_name": "Firm", "client_name": "Client's name",
            "recipient_name": "Recipient's name", "recipient_address": "Recipient's address",
            "subject": "Matter / subject", "response_days": "Days to respond (e.g. 14)",
        },
        {"facts": FACTS, "legal_grounds": "The basis of the claim.", "demands": "What the client requires, including any sum claimed and how it is calculated."},
        [
            (None, ["{{date}}", "{{recipient_name}}", "{{recipient_address}}", "Dear {{recipient_name}},", "Re: {{subject}}",
                    "We act for {{client_name}}. We write in accordance with the Practice Direction on Pre-Action Conduct and Protocols."]),
            ("Background", ["{{ai:facts}}"]),
            ("Basis of claim", ["{{ai:legal_grounds}}"]),
            ("What our client requires", ["{{ai:demands}}"]),
            ("Response", [
                "Please respond within {{response_days}} days of the date of this letter, with any documents on which you rely. "
                "If you do not respond, our client may issue proceedings without further notice and will refer the court to "
                "your failure to comply with the Practice Direction when seeking costs.",
                "If you are not legally represented, we recommend that you seek independent legal advice.",
            ]),
        ],
        ["Yours faithfully,", "{{solicitor_name}}", "{{firm_name}}"],
    ),
    # ------------------------------------------------------------------ bail
    Template(
        "bail_application", "IN", "Bail application (India, BNSS)", "IN THE COURT OF {{court_name}}",
        {
            "court_name": "Court, e.g. 'the Sessions Judge, Pune'", "case_reference": "Bail application number, if allotted",
            "applicant_name": "Accused/applicant's name", "respondent": "Respondent (usually 'State of ...')",
            "fir_number": "FIR number and year", "police_station": "Police station",
            "sections_charged": "Sections charged (e.g. '318(4) BNS')",
            "bail_type": "REGULAR BAIL or ANTICIPATORY BAIL",
            "bail_section": "BNSS section: 480 or 483 (regular), 482 (anticipatory)",
            "advocate_name": "Advocate's name", "place": "Place",
        },
        {
            "facts": FACTS + " Include the date of arrest or apprehension of arrest, and custody status.",
            "grounds": "Numbered grounds for bail: nature of accusation, parity, custody period, no flight risk, cooperation, "
                       "roots in society, health or family circumstances, triple test, and applicable precedent if provided.",
            "prayer": "The precise relief sought, e.g. 'release the applicant on bail in FIR No. ... on such terms as this Hon'ble Court deems fit'.",
        },
        [
            (None, ["{{case_reference}}", "IN THE MATTER OF:", "{{applicant_name}} ... Applicant/Accused", "Versus",
                    "{{respondent}} ... Respondent",
                    "FIR No. {{fir_number}}, Police Station {{police_station}}, under Sections {{sections_charged}}"]),
            ("APPLICATION FOR {{bail_type}} UNDER SECTION {{bail_section}} OF THE BHARATIYA NAGARIK SURAKSHA SANHITA, 2023",
             ["MOST RESPECTFULLY SHOWETH:", "{{ai:facts}}"]),
            ("Grounds", ["{{ai:grounds}}"]),
            ("Prayer", [
                "In view of the facts and circumstances stated above, it is most respectfully prayed that this Hon'ble Court may be pleased to:",
                "{{ai:prayer}}",
                "And pass any other order(s) as this Hon'ble Court may deem fit and proper in the interest of justice.",
            ]),
        ],
        ["Applicant/Accused", "Through Counsel", "{{advocate_name}}", "Place: {{place}}", "Date: {{date}}"],
    ),
    Template(
        "bail_application", "*", "Application for bail / pretrial release", "APPLICATION FOR BAIL",
        {
            "court_name": "Court", "case_number": "Case number", "prosecution": "Prosecuting authority (e.g. 'The People', 'R', 'United States')",
            "defendant_name": "Defendant's name", "charges": "Charges", "counsel_name": "Counsel's name",
        },
        {"facts": FACTS, "grounds": "Numbered grounds supporting release, addressing the statutory factors for the jurisdiction.",
         "prayer": "The release and conditions sought."},
        [
            (None, ["{{court_name}}", "{{prosecution}} v. {{defendant_name}}", "Case No. {{case_number}}", "Charges: {{charges}}"]),
            ("Background", ["{{ai:facts}}"]),
            ("Grounds", ["{{ai:grounds}}"]),
            ("Relief sought", ["For the foregoing reasons, {{defendant_name}} respectfully asks the Court to:", "{{ai:prayer}}"]),
        ],
        ["Respectfully submitted,", "{{counsel_name}}", "Counsel for {{defendant_name}}", "{{date}}"],
    ),
    # ------------------------------------------------------------------ cease and desist
    Template(
        "cease_and_desist", "*", "Cease and desist letter", "CEASE AND DESIST",
        {
            "sender_name": "Lawyer or firm sending the letter", "client_name": "Client's name",
            "recipient_name": "Recipient's name", "recipient_address": "Recipient's address",
            "subject": "Subject (e.g. 'Infringement of the BRIGHTPATH trade mark')", "response_deadline": "Deadline to confirm compliance",
        },
        {"facts": "Numbered paragraphs describing the client's rights and the recipient's conduct.",
         "legal_grounds": "Why the conduct is unlawful, citing applicable law.",
         "demands": "Numbered list of what the recipient must stop and do (e.g. cease use, remove material, account, undertake)."},
        [
            (None, ["{{date}}", "To: {{recipient_name}}, {{recipient_address}}", "Re: {{subject}}",
                    "We act for {{client_name}}. This letter requires you to cease and desist from the conduct described below."]),
            ("The conduct", ["{{ai:facts}}"]),
            ("Our client's rights", ["{{ai:legal_grounds}}"]),
            ("What you must do", [
                "{{ai:demands}}",
                "Please confirm in writing by {{response_deadline}} that you have complied. If you do not, our client may take "
                "legal action without further notice, including seeking an injunction, damages and costs. This letter is "
                "without prejudice to all of our client's rights and remedies, which are expressly reserved.",
            ]),
        ],
        ["{{sender_name}}", "For and on behalf of {{client_name}}"],
    ),
    # ------------------------------------------------------------------ NDA
    Template(
        "nda", "*", "Mutual non-disclosure agreement", "MUTUAL NON-DISCLOSURE AGREEMENT",
        {
            "party_a_name": "First party's name", "party_a_details": "First party's description and address",
            "party_b_name": "Second party's name", "party_b_details": "Second party's description and address",
            "purpose": "Purpose of sharing information", "term_years": "Confidentiality period in years",
            "governing_law": "Governing law", "courts": "Courts with jurisdiction",
        },
        {"additional_terms": "Any additional clauses the context specifically requires (e.g. non-solicitation), as numbered "
                             "clauses starting at 9; or exactly 'None'."},
        [
            (None, ["This Mutual Non-Disclosure Agreement is made on {{date}} between {{party_a_name}}, {{party_a_details}}, "
                    "and {{party_b_name}}, {{party_b_details}} (each a \"Party\").",
                    "The Parties wish to exchange Confidential Information for the purpose of {{purpose}} (the \"Purpose\")."]),
            ("1. Confidential Information", ["\"Confidential Information\" means all information disclosed by a Party (the "
                                              "\"Discloser\") to the other (the \"Recipient\") in connection with the Purpose, in any form, "
                                              "that is marked confidential or would reasonably be understood to be confidential."]),
            ("2. Obligations", ["The Recipient shall (a) use Confidential Information only for the Purpose; (b) not disclose it "
                                "except to its employees and advisers who need to know it for the Purpose and are bound by "
                                "equivalent obligations; and (c) protect it with at least reasonable care."]),
            ("3. Exclusions", ["These obligations do not apply to information that (a) is or becomes public other than through "
                               "the Recipient's breach; (b) was lawfully known to the Recipient before disclosure; (c) is lawfully "
                               "received from a third party without restriction; or (d) is independently developed. The Recipient "
                               "may disclose information when required by law, giving the Discloser prompt notice where lawful."]),
            ("4. Term", ["This Agreement applies to information disclosed within two years of its date, and the obligations "
                         "continue for {{term_years}} years after each disclosure."]),
            ("5. Return of information", ["On request, the Recipient shall promptly return or destroy the Discloser's "
                                          "Confidential Information, except as required to be retained by law."]),
            ("6. No licence", ["No licence or other right is granted under any intellectual property by this Agreement."]),
            ("7. Remedies", ["Damages may not be an adequate remedy for breach, and the Discloser may seek injunctive relief in "
                             "addition to any other remedy."]),
            ("8. Governing law", ["This Agreement is governed by {{governing_law}}, and the courts of {{courts}} have "
                                  "exclusive jurisdiction."]),
            (None, ["{{ai:additional_terms}}"]),
        ],
        ["Signed for {{party_a_name}}: ____________________", "Signed for {{party_b_name}}: ____________________"],
    ),
    # ------------------------------------------------------------------ affidavits
    Template(
        "affidavit", "IN", "Affidavit (India)", "AFFIDAVIT",
        {
            "deponent_name": "Deponent's name", "deponent_relation": "e.g. 'son of ...' / 'daughter of ...'",
            "deponent_age": "Deponent's age", "deponent_address": "Deponent's address", "place": "Place of verification",
        },
        {"statements": "Numbered paragraphs in the first person, one fact per paragraph, only facts within the deponent's knowledge."},
        [
            (None, ["I, {{deponent_name}}, {{deponent_relation}}, aged about {{deponent_age}} years, residing at "
                    "{{deponent_address}}, do hereby solemnly affirm and declare as under:"]),
            (None, ["{{ai:statements}}", "DEPONENT"]),
            ("VERIFICATION", ["Verified at {{place}} on {{date}} that the contents of the above affidavit are true and correct to "
                              "the best of my knowledge and belief, and nothing material has been concealed therefrom."]),
        ],
        ["DEPONENT"],
    ),
    Template(
        "affidavit", "US", "Affidavit (US)", "AFFIDAVIT OF {{deponent_name}}",
        {"deponent_name": "Affiant's name", "state": "State", "county": "County"},
        {"statements": "Numbered paragraphs in the first person, one fact per paragraph, within the affiant's personal knowledge."},
        [
            (None, ["STATE OF {{state}}", "COUNTY OF {{county}}",
                    "I, {{deponent_name}}, being first duly sworn, depose and state as follows:"]),
            (None, ["{{ai:statements}}",
                    "I declare under penalty of perjury under the laws of the State of {{state}} that the foregoing is true and correct."]),
        ],
        ["{{deponent_name}}", "Subscribed and sworn to before me on {{date}}.", "____________________", "Notary Public"],
    ),
    Template(
        "affidavit", "UK", "Affidavit (England & Wales)", "AFFIDAVIT",
        {
            "court_name": "Court (e.g. 'HIGH COURT OF JUSTICE, KING'S BENCH DIVISION')", "claim_number": "Claim number",
            "party": "Party on whose behalf it is made", "deponent_name": "Deponent's name",
            "deponent_address": "Deponent's address", "deponent_occupation": "Deponent's occupation", "place": "Place sworn",
        },
        {"statements": "Numbered paragraphs in the first person, indicating which facts are from the deponent's own knowledge "
                       "and the source of any others."},
        [
            (None, ["Party: {{party}} | Deponent: {{deponent_name}} | No. 1 | Date sworn: {{date}}",
                    "IN THE {{court_name}}", "Claim No. {{claim_number}}",
                    "I, {{deponent_name}}, of {{deponent_address}}, {{deponent_occupation}}, MAKE OATH and say as follows:"]),
            (None, ["{{ai:statements}}"]),
        ],
        ["Sworn at {{place}} on {{date}}", "Before me, ____________________", "A solicitor / Commissioner for Oaths"],
    ),
]


class DraftingError(ValueError):
    pass


def template_family(jurisdiction: JurisdictionProfile | None) -> str:
    if jurisdiction is None:
        return "*"
    return "US" if jurisdiction.code.startswith("US") else jurisdiction.code


def get_template(doc_type: str, jurisdiction: JurisdictionProfile | None) -> Template:
    if doc_type not in DOC_TYPES:
        raise DraftingError(f"Unknown document type {doc_type!r}. Choose one of: {', '.join(DOC_TYPES)}")
    family = template_family(jurisdiction)
    candidates = [t for t in TEMPLATES if t.doc_type == doc_type]
    for wanted in (family, "*"):
        for t in candidates:
            if t.family == wanted:
                return t
    if jurisdiction is None:
        return candidates[0]
    raise DraftingError(f"No {DOC_TYPES[doc_type]} template for {jurisdiction.name}")


def automatic_variables(jurisdiction: JurisdictionProfile | None, today: date | None = None) -> dict[str, str]:
    today = today or date.today()
    formatted = today.strftime(jurisdiction.date_format if jurisdiction else "%d %B %Y")
    values = {"date": formatted.lstrip("0").replace(" 0", " ")}  # "5 March 2026", "March 5, 2026"
    if jurisdiction is not None:
        if jurisdiction.code == "IN":
            values.update(governing_law="the laws of India")
        elif jurisdiction.code == "UK":
            values.update(governing_law="the laws of England and Wales", courts="England and Wales")
        elif jurisdiction.code.startswith("US-"):
            state = US_STATES[jurisdiction.code[3:]]
            values.update(governing_law=f"the laws of the State of {state}", state=state, courts=f"the State of {state}")
    return values


# ----------------------------------------------------------------------------- rendering


def _placeholder(description: str) -> str:
    return f"[●{description}●]"


def render_text(text: str, variables: dict[str, str], slots: dict[str, str], template: Template, missing: set[str]) -> str:
    def replace(match: re.Match) -> str:
        is_slot, name = bool(match.group(1)), match.group(2)
        if is_slot:
            value = slots.get(name, "").strip()
            if not value:
                missing.add(f"Text to draft: {template.slots.get(name, name)}")
                return _placeholder(f"to be drafted: {name.replace('_', ' ')}")
            return value
        value = (variables.get(name) or "").strip()
        if not value:
            description = template.variables.get(name, name.replace("_", " "))
            missing.add(description)
            return _placeholder(description)
        return value

    return _FIELD.sub(replace, text)


def render(template: Template, variables: dict[str, str], slots: dict[str, str]) -> dict:
    """Fill the template. Slot passages that are a whole paragraph become several paragraphs."""
    missing: set[str] = set()
    sections = []
    for heading, paragraphs in template.sections:
        rendered: list[str] = []
        for paragraph in paragraphs:
            whole_slot = _FIELD.fullmatch(paragraph.strip())
            if whole_slot and whole_slot.group(1):
                value = slots.get(whole_slot.group(2), "").strip()
                if value.lower() in ("none", "none."):
                    continue  # the template marks this passage as optional
            text = render_text(paragraph, variables, slots, template, missing)
            if whole_slot and whole_slot.group(1):
                rendered += [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
            else:
                rendered.append(text)
        if rendered:
            sections.append(
                {"heading": render_text(heading, variables, slots, template, missing) if heading else None, "paragraphs": rendered}
            )
    return {
        "title": render_text(template.title, variables, slots, template, missing),
        "sections": sections,
        "signature_block": [render_text(line, variables, slots, template, missing) for line in template.signature],
        "missing": sorted(missing),
    }


def template_spec(template: Template, provided: dict[str, str]) -> str:
    """How the template is described to Claude."""
    lines = [f"Document: {template.name}", "", "Variables (fill only from the case context or the lawyer's instructions):"]
    for name, description in template.variables.items():
        if name in provided:
            lines.append(f"- {name}: already provided by the lawyer as {provided[name]!r}")
        else:
            lines.append(f"- {name}: {description}")
    lines += ["", "Passages to write (slots):"] + [f"- {n}: {d}" for n, d in template.slots.items()]
    lines += ["", "Template wording (for context):"]
    for heading, paragraphs in template.sections:
        if heading:
            lines.append(f"## {heading}")
        lines += paragraphs
    return "\n".join(lines)


# ----------------------------------------------------------------------------- export


def _segments(text: str) -> list[tuple[str, bool]]:
    """Split text into (piece, is_placeholder) runs."""
    out, cursor = [], 0
    for match in PLACEHOLDER.finditer(text):
        if match.start() > cursor:
            out.append((text[cursor:match.start()], False))
        out.append((match.group(0), True))
        cursor = match.end()
    if cursor < len(text):
        out.append((text[cursor:], False))
    return out


UNREVIEWED_MARK = "AI DRAFT · UNREVIEWED"
UNREVIEWED_NOTE = "Not yet approved by a lawyer. Review every fact, date, amount and citation before use."


def unreviewed(draft: dict) -> bool:
    return not draft.get("approved_at")


def to_docx(draft: dict) -> bytes:
    """An editable Word document; unfilled placeholders are highlighted yellow.
    Unapproved drafts carry an "AI draft · unreviewed" header on every page."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    def add_runs(paragraph, text: str, bold: bool = False):
        for piece, is_placeholder in _segments(text):
            run = paragraph.add_run(piece)
            run.bold = bold
            if is_placeholder:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    if unreviewed(draft):
        from docx.shared import RGBColor

        header = document.sections[0].header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mark = header.add_run(f"{UNREVIEWED_MARK} · {UNREVIEWED_NOTE}")
        mark.bold, mark.font.size, mark.font.color.rgb = True, Pt(9), RGBColor(0xB4, 0x1C, 0x1C)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(title, draft["title"], bold=True)

    for section in draft["content"]["sections"]:
        if section.get("heading"):
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(12)
            add_runs(heading, section["heading"], bold=True)
        for text in section["paragraphs"]:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(paragraph, text)

    signature = draft["content"].get("signature_block") or []
    if signature:
        document.add_paragraph()
        for line in signature:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_runs(paragraph, line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_pdf(draft: dict) -> bytes:
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    from app.services.pdf_service import _get_fonts

    regular, bold = _get_fonts()
    base = getSampleStyleSheet()["Normal"]
    body = ParagraphStyle("body", parent=base, fontName=regular, fontSize=11, leading=16, alignment=TA_JUSTIFY, spaceAfter=7)
    title = ParagraphStyle("title", parent=body, fontName=bold, fontSize=13, alignment=TA_CENTER, spaceAfter=14)
    heading = ParagraphStyle("heading", parent=body, fontName=bold, spaceBefore=8)
    signature = ParagraphStyle("signature", parent=body, alignment=TA_RIGHT, spaceAfter=2)

    def markup(text: str) -> str:
        return "".join(
            f'<font backColor="#FFF176">{escape(piece)}</font>' if is_placeholder else escape(piece)
            for piece, is_placeholder in _segments(text)
        ).replace("\n", "<br/>")

    story = [Paragraph(markup(draft["title"]), title)]
    for section in draft["content"]["sections"]:
        if section.get("heading"):
            story.append(Paragraph(markup(section["heading"]), heading))
        story += [Paragraph(markup(p), body) for p in section["paragraphs"]]
    if draft["content"].get("signature_block"):
        story.append(Spacer(1, 10 * mm))
        story += [Paragraph(markup(line), signature) for line in draft["content"]["signature_block"]]

    def watermark(canvas, doc):
        if not unreviewed(draft):
            return
        width, height = A4
        canvas.saveState()
        canvas.setFont(bold, 9)
        canvas.setFillColorRGB(0.71, 0.11, 0.11)
        canvas.drawCentredString(width / 2, height - 14 * mm, f"{UNREVIEWED_MARK} · {UNREVIEWED_NOTE}")
        canvas.setFont(bold, 54)
        canvas.setFillColorRGB(0.85, 0.2, 0.2, alpha=0.12)
        canvas.translate(width / 2, height / 2)
        canvas.rotate(40)
        canvas.drawCentredString(0, 0, UNREVIEWED_MARK)
        canvas.restoreState()

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=25 * mm, rightMargin=25 * mm, topMargin=25 * mm,
                      bottomMargin=25 * mm, title=draft["title"]).build(story, onFirstPage=watermark, onLaterPages=watermark)
    return buffer.getvalue()


def export_filename(draft: dict, extension: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", draft.get("title") or "draft").strip("-").lower()[:60] or "draft"
    return f"{slug}.{extension}"
