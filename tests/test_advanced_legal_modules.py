"""Phase 7: contract redlining, jurisdictions, contradiction analysis, drafting."""

import asyncio
import io
import json
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.dependencies.auth import CurrentUser, get_current_user
from app.main import app
from app.models.analysis import (
    ClauseRevision,
    ClauseRiskAssessment,
    ContractDiffAnalysis,
    ContradictionAnalysis,
    ContradictionFinding,
    DraftFill,
    NamedValue,
)
from app.routers import audio_stream
from app.services import (
    anthropic_service,
    contract_diff_service as diff,
    document_service,
    drafting_engine,
    evidence_service as evidence,
    jurisdiction_service as juris,
    memory_service,
    prompt_engine,
    rag_service,
    transcript_service,
)
from app.services.contract_diff_service import align, split_contract
from app.services.evidence_service import EvidenceSource, StatementLine
from tests.conftest import ORG_ID

USER = CurrentUser(id="11111111-1111-1111-1111-111111111111", email="lawyer@acme.law", access_token="good-token")
SSN = "536-22-8172"

ORIGINAL = """1. Term. This Agreement lasts 12 months from the Effective Date.

2. Payment. The Client shall pay each invoice within 30 days of receipt.

3. Liability. Each party's total liability is capped at the fees paid in the prior 12 months.

4. Confidentiality. Each party shall keep the other party's information confidential.

5. Governing law. This Agreement is governed by the laws of India."""

REVISED = """1. Term. This Agreement lasts 12 months from the Effective Date.

2. Payment. The Client shall pay each invoice within 15 days of receipt, with interest at 2% per month on late payments.

3. Confidentiality. Each party shall keep the other party's information confidential.

4. Indemnity. The Client shall indemnify the Supplier against all claims arising from the Services.

5. Governing law. This Agreement is governed by the laws of India."""


@pytest.fixture
def client():
    app.dependency_overrides[get_current_user] = lambda: USER
    yield TestClient(app)
    app.dependency_overrides.clear()


# ============================================================================ contract redlining


def test_clauses_are_aligned_across_versions():
    changes = align(split_contract(ORIGINAL), split_contract(REVISED))
    assert [(c.status, c.label) for c in changes] == [
        ("unchanged", "1"), ("modified", "2"), ("deleted", "3"), ("renumbered", "3"), ("added", "4"), ("unchanged", "5"),
    ]


def test_inline_diff_reconstructs_both_versions():
    modified = next(c for c in align(split_contract(ORIGINAL), split_contract(REVISED)) if c.status == "modified")
    ops = modified.inline_diff
    assert "".join(o["text"] for o in ops if o["op"] != "insert") == modified.original.text
    assert "".join(o["text"] for o in ops if o["op"] != "delete") == modified.revised.text
    assert {"op": "delete", "text": "30"} in ops and any(o["op"] == "insert" and "15" in o["text"] for o in ops)


def test_page_markers_from_extraction_are_ignored():
    paged = "[Page 1]\n" + ORIGINAL.replace("3. Liability", "[Page 2]\n3. Liability")
    assert [c.status for c in align(split_contract(paged), split_contract(ORIGINAL))] == ["unchanged"] * 5


def assessment(clause_id, impact, severity, counter=None):
    return ClauseRiskAssessment(clause_id=clause_id, change_summary="x", risk_impact=impact, severity=severity,
                                recommended_counter_clause=counter)


@pytest.mark.parametrize(
    "assessments, expected",
    [
        ([("increased", "high"), ("decreased", "low")], ("Increased", 2)),
        ([("increased", "low"), ("decreased", "medium")], ("Decreased", -1)),
        ([("increased", "medium"), ("decreased", "medium"), ("neutral", "high")], ("Neutral", 0)),
        ([], ("Neutral", 0)),
    ],
)
def test_risk_delta_weights_severity(assessments, expected):
    assert diff.risk_delta([assessment(f"c{i}", imp, sev) for i, (imp, sev) in enumerate(assessments)]) == expected


@pytest.fixture
def fake_redline(monkeypatch):
    seen = []

    async def analyze(changes_xml, *, jurisdiction=None, our_side=None):
        seen.append({"xml": changes_xml, "jurisdiction": jurisdiction, "our_side": our_side})
        ids = [line.split('"')[1] for line in changes_xml.splitlines() if line.startswith("<clause id=")]
        statuses = {line.split('"')[1]: line.split('status="')[1].split('"')[0] for line in changes_xml.splitlines() if line.startswith("<clause id=")}
        clauses = []
        for cid in ids:
            if statuses[cid] == "added":
                clauses.append(assessment(cid, "increased", "high", "4. Indemnity. Each party shall indemnify the other for losses caused by its breach, capped as in clause 3. Notices to [EMAIL]."))
            elif statuses[cid] == "deleted":
                clauses.append(assessment(cid, "increased", "high"))
            else:
                clauses.append(assessment(cid, "increased", "medium"))
        return ContractDiffAnalysis(clauses=clauses, overall_summary="The revision shifts risk to the Client.",
                                    key_concerns=["Liability cap removed"])

    monkeypatch.setattr(anthropic_service, "is_configured", lambda: True)
    monkeypatch.setattr(anthropic_service, "analyze_contract_changes", analyze)
    return seen


def test_compare_sends_only_changed_clauses_masked_and_computes_the_delta(fake_redline):
    revised = REVISED.replace("against all claims", f"against all claims (notices to legal@supplier.com, SSN {SSN})")
    report = asyncio.run(diff.compare(ORIGINAL, revised, jurisdiction=juris.get_profile("IN"), our_side="the Client"))

    sent = fake_redline[0]["xml"]
    assert sent.count("<clause id=") == 3  # modified, deleted, added: not the unchanged or renumbered ones
    assert "legal@supplier.com" not in sent and SSN not in sent and "[EMAIL]" in sent
    assert fake_redline[0]["jurisdiction"].code == "IN" and fake_redline[0]["our_side"] == "the Client"

    assert report["risk_delta"] == "Increased" and report["risk_score"] == 3 + 3 + 2
    assert [c["label"] for c in report["additions"]] == ["4"] and [c["label"] for c in report["deletions"]] == ["3"]
    counter = report["recommended_counter_clauses"][0]
    assert counter["label"] == "4" and "legal@supplier.com" in counter["text"]  # unmasked for the lawyer
    assert report["stats"] == {"clauses_original": 5, "clauses_revised": 5, "unchanged": 2, "renumbered": 1,
                               "modified": 1, "added": 1, "deleted": 1, "analyzed": 3}


def test_compare_without_analysis_still_redlines(monkeypatch):
    monkeypatch.setattr(anthropic_service, "is_configured", lambda: False)
    report = asyncio.run(diff.compare(ORIGINAL, REVISED))
    assert report["risk_delta"] == "Not analyzed" and len(report["modifications"]) == 1


def test_compare_endpoint(client, fake_redline, audit_log, monkeypatch):
    docs = {"aaaaaaaa-0000-0000-0000-000000000001": ORIGINAL, "aaaaaaaa-0000-0000-0000-000000000002": REVISED}
    monkeypatch.setattr(document_service, "get_document", lambda token, did, include_text=False, org_id=None:
                        {"id": did, "filename": f"{did[-1]}.pdf", "extracted_text": docs[did]} if did in docs else None)
    ids = list(docs)
    response = client.post("/api/v1/documents/compare", json={"original_document_id": ids[0], "revised_document_id": ids[1], "jurisdiction": "india"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["original"]["filename"] == "1.pdf" and body["risk_delta"] == "Increased" and body["jurisdiction"] == "IN"
    assert audit_log[-1]["action"] == "COMPARE_DOCUMENTS"

    assert client.post("/api/v1/documents/compare", json={"original_document_id": ids[0], "revised_document_id": ids[0]}).status_code == 400
    assert client.post("/api/v1/documents/compare", json={"original_document_id": ids[0], "revised_document_id": ids[1], "jurisdiction": "Narnia"}).status_code == 422
    assert client.post("/api/v1/documents/compare", json={"original_document_id": ids[0], "revised_document_id": str(uuid4())}).status_code == 404


# ============================================================================ jurisdictions


@pytest.mark.parametrize(
    "citation, style, expected",
    [
        (juris.CaseCitation("Brown v Board of Education", 1954, "U.S.", 483, 347, pinpoint=495), "bluebook",
         "Brown v. Board of Education, 347 U.S. 483, 495 (1954)"),
        (juris.CaseCitation("Smith vs. Jones", 1999, "F.3d", 456, 123, court="9th Cir."), "bluebook", "Smith v. Jones, 123 F.3d 456 (9th Cir. 1999)"),
        (juris.CaseCitation("Donoghue v. Stevenson", 1932, "AC", 562, court="HL", pinpoint=580), "oscola", "Donoghue v Stevenson [1932] AC 562 (HL) 580"),
        (juris.CaseCitation("R (Miller) v Prime Minister", 2019, neutral="[2019] UKSC 41"), "oscola", "R (Miller) v Prime Minister [2019] UKSC 41"),
        (juris.CaseCitation("State of Maharashtra v Ravi Kumar", 2019, "SCC", 123, 5), "indian", "State of Maharashtra v. Ravi Kumar, (2019) 5 SCC 123"),
        (juris.CaseCitation("Kesavananda Bharati v State of Kerala", 1973, "AIR SC", 1461), "indian", "Kesavananda Bharati v. State of Kerala, AIR 1973 SC 1461"),
    ],
)
def test_case_citations_follow_regional_styles(citation, style, expected):
    assert juris.format_case(citation, style) == expected


def test_statute_citations_and_new_indian_codes():
    assert juris.format_statute(juris.StatuteCitation("U.S.C.", "1983", title_number=42), "bluebook") == "42 U.S.C. § 1983"
    assert juris.format_statute(juris.StatuteCitation("Companies Act", "172", 2006), "oscola") == "Companies Act 2006, s 172"
    assert juris.format_statute(juris.StatuteCitation("Bharatiya Nyaya Sanhita", "103", 2023), "indian") == "Section 103, Bharatiya Nyaya Sanhita, 2023"
    assert juris.new_code_section("IPC", "302") == "Section 103(1) BNS (formerly Section 302 IPC)"
    assert juris.new_code_section("CrPC", "438") == "Section 482 BNSS (formerly Section 438 CrPC)"
    assert juris.new_code_section("Evidence Act", "65b") == "Section 63 BSA (formerly Section 65B Evidence Act)"
    assert juris.new_code_section("IPC", "9999") is None


def test_jurisdiction_codes_are_normalized():
    assert [juris.normalize(v) for v in ["india", "England & Wales", "California", "us-ny", "USA", "Mars", ""]] == \
        ["IN", "UK", "US-CA", "US-NY", "US", None, None]
    assert juris.rag_filter("US-CA") == ["US-CA", "US"] and juris.rag_filter("IN") == ["IN"] and juris.rag_filter(None) is None
    with pytest.raises(juris.UnknownJurisdiction):
        juris.require_profile("Atlantis")


def test_jurisdiction_rules_are_injected_into_the_cached_prompt():
    system = prompt_engine.build_system("BASE", jurisdiction=juris.get_profile("IN"))
    text = system[0]["text"]
    assert system[0]["cache_control"] == {"type": "ephemeral"}
    assert "Bharatiya Nyaya Sanhita, 2023 (BNS)" in text and "1 July 2024" in text
    assert "Citation style: Indian citation (SCC / AIR / neutral citations)" in text
    uk = prompt_engine.build_system("BASE", jurisdiction=juris.get_profile("UK"))[0]["text"]
    assert "OSCOLA" in uk and "<Party> v <Party> [<year>] UKSC <number>" in uk


def test_rag_search_is_filtered_by_jurisdiction():
    calls = []
    db = SimpleNamespace(rpc=lambda fn, params: calls.append(params) or SimpleNamespace(execute=lambda: SimpleNamespace(data=[])))
    embedder = SimpleNamespace(embed_query=lambda text: [0.1, 0.2])
    rag_service.retrieve_relevant_precedents(ORG_ID, "The landlord kept my security deposit after I left.", db=db,
                                             embedder=embedder, jurisdictions=juris.rag_filter("US-CA"))
    assert calls[0]["filter_jurisdictions"] == ["US-CA", "US"]


def test_knowledge_search_endpoint_applies_the_filter(client, monkeypatch):
    seen = {}

    async def retrieve(org_id, query, top_k, *, db, match_threshold=None, jurisdictions=None):
        seen["jurisdictions"] = jurisdictions
        return []

    monkeypatch.setattr(rag_service, "aretrieve_relevant_precedents", retrieve)
    assert client.post("/api/v1/knowledge/search", json={"query": "deposit return", "jurisdiction": "UK"}).status_code == 200
    assert seen["jurisdictions"] == ["UK"]
    assert client.post("/api/v1/knowledge/search", json={"query": "deposit return", "jurisdiction": "Mars"}).status_code == 422


def test_live_session_runs_under_the_chosen_jurisdiction(monkeypatch, fake_agents):
    seen = {"retrieve": [], "saved": None}

    async def retrieve(org_id, text, top_k, *, db, jurisdictions=None):
        seen["retrieve"].append(jurisdictions)
        return []

    async def summarize(transcript, analysis, **context):
        seen["summary_jurisdiction"] = context.get("jurisdiction")
        from app.models.analysis import ConsultationSummary
        return ConsultationSummary(title="t", summary="s")

    monkeypatch.setattr(audio_stream, "verify_access_token", lambda token: USER)
    monkeypatch.setattr(rag_service, "aretrieve_relevant_precedents", retrieve)
    monkeypatch.setattr(anthropic_service, "summarize_consultation", summarize)
    monkeypatch.setattr(transcript_service, "save_live_session", lambda token, uid, row: seen.update(saved=row) or {"id": "t1", **row})
    monkeypatch.setattr("app.services.n8n_service.is_configured", lambda: False)

    with TestClient(app).websocket_connect("/api/v1/ws/audio?token=good-token&jurisdiction=california") as ws:
        assert ws.receive_json()["jurisdiction"] == "US-CA"
        ws.send_json({"type": "transcript_chunk", "speaker": "Client", "text": "My landlord kept my deposit."})
        while ws.receive_json()["type"] != "analysis":
            pass
        ws.send_json({"type": "end"})
        while ws.receive_json()["type"] != "session_ended":
            pass

    assert fake_agents.calls["intake"][0]["jurisdiction"].code == "US-CA"
    assert fake_agents.calls["risk"][0]["jurisdiction"].code == "US-CA"
    assert seen["retrieve"] == [["US-CA", "US"]] and seen["summary_jurisdiction"].code == "US-CA"
    segments = seen["saved"]["metadata"]["segments"]
    assert segments[0]["speaker"] == "Client" and segments[0]["offset"] >= 0
    assert seen["saved"]["metadata"]["jurisdiction"] == "US-CA"


# ============================================================================ contradictions

DEPOSITION = """[00:01:05] Witness: I left the office at 6 PM that evening.
[00:01:20] Witness: The driver, Mr Rao, was at the warehouse in Thane at 9 PM.
[00:02:10] Witness: I never signed the delivery receipt.
[00:03:00] Witness: You can reach me on 212-555-0187."""

FIR = "FIR No. 112/2026, PS Thane. Complainant states the vehicle MH-04-1234 left the Thane warehouse at 7:30 PM driven by Mr Rao and was stopped at Vashi toll plaza at 9:05 PM."
RECEIPT = "Delivery receipt 4471: received in good condition. Signed: A. Kulkarni (witness), 6:40 PM."


def test_statement_lines_keep_timestamps_speakers_and_pages():
    statement = evidence.parse_statement_text("[Page 2]\n(00:10) Q: Where were you?\n00:12 Witness: At home.\n\nPlain line")
    assert [(l.number, l.timestamp, l.speaker, l.location, l.text) for l in statement.lines] == [
        (1, "00:10", "Q", "Page 2", "Where were you?"), (2, "00:12", "Witness", "Page 2", "At home."), (3, None, None, "Page 2", "Plain line"),
    ]


def test_saved_transcripts_get_timestamps_from_session_offsets():
    row = {"id": "t1", "title": "Deposition", "content": "Witness: Hello.\nCounsel: Hi.",
           "metadata": {"segments": [{"offset": 3.2}, {"offset": 3725}]}}
    statement = evidence.statement_from_transcript(row)
    assert [l.timestamp for l in statement.lines] == ["00:00:03", "01:02:05"]


def test_deepgram_utterances_become_timestamped_lines():
    text = evidence.transcript_text_from_utterances([
        {"start": 65.4, "speaker": 0, "transcript": "I left at six."}, {"start": 70, "speaker": 1, "transcript": "Are you sure?"},
        {"start": 71, "speaker": 1, "transcript": "  "},
    ])
    assert text == "[00:01:05] Speaker 1: I left at six.\n[00:01:10] Speaker 2: Are you sure?"


def finding(line, quote, source, source_quote, severity="high", kind="temporal"):
    return ContradictionFinding(statement_line=line, statement_quote=quote, source_id=source, source_quote=source_quote,
                                kind=kind, severity=severity, explanation="Conflicts.", cross_examination_questions=["Isn't it true...?"])


def test_findings_are_verified_against_the_actual_texts():
    lines = {1: "I left the office at 6 PM that evening.", 2: "The driver was at the warehouse in Thane at 9 PM."}
    sources = {"S1": EvidenceSource("S1", "document", "FIR", FIR)}
    kept, discarded = evidence.verify_findings([
        finding(2, "at the warehouse in Thane at 9 PM", "S1", "stopped at Vashi toll plaza at 9:05 PM"),
        finding(1, "at the warehouse in Thane at 9 PM", "S1", "left the Thane warehouse at 7:30 PM"),  # wrong line: corrected
        finding(2, "the warehouse in Thane", "S1", "the vehicle was in Pune all night"),  # invented source quote
        finding(1, "I was asleep at home", "S1", "left the Thane warehouse at 7:30 PM"),  # invented statement quote
        finding(1, "I left the office", "S9", "anything"),  # unknown source
    ], lines, sources)
    assert [(f.statement_line, f.source_quote) for f in kept] == [
        (2, "stopped at Vashi toll plaza at 9:05 PM"), (2, "left the Thane warehouse at 7:30 PM")]
    assert discarded == 3


@pytest.fixture
def fake_contradictions(monkeypatch):
    seen = {}

    async def analyze(statement_xml, sources_xml, *, jurisdiction=None):
        seen.update(statement=statement_xml, sources=sources_xml, jurisdiction=jurisdiction)
        return ContradictionAnalysis(
            contradictions=[
                finding(2, "was at the warehouse in Thane at 9 PM", "S1", "stopped at Vashi toll plaza at 9:05 PM"),
                finding(3, "I never signed the delivery receipt", "S2", "Signed: A. Kulkarni (witness)", severity="medium", kind="direct"),
                finding(1, "I left the office at 6 PM", "S2", "Signed at 11 PM in Pune", severity="high"),  # fabricated
            ],
            consistent_points=["The witness agrees Mr Rao drove the vehicle."],
            overall_assessment="Two material inconsistencies.",
        )

    monkeypatch.setattr(anthropic_service, "is_configured", lambda: True)
    monkeypatch.setattr(anthropic_service, "analyze_contradictions", analyze)
    return seen


def test_contradiction_report_is_timestamped_ranked_and_verified(fake_contradictions):
    statement = evidence.parse_statement_text(DEPOSITION, "Deposition of A. Kulkarni")
    sources = [EvidenceSource("S1", "document", "FIR 112/2026", FIR, "doc-fir"), EvidenceSource("S2", "document", "Delivery receipt", RECEIPT, "doc-rcpt")]
    report = asyncio.run(evidence.analyze(statement, sources, jurisdiction=juris.get_profile("IN")))

    assert [(c["statement_line"], c["timestamp"], c["severity_score"]) for c in report["contradictions"]] == [
        (2, "00:01:20", 0.9), (3, "00:02:10", 0.6)]
    top = report["contradictions"][0]
    assert top["source"] == {"id": "S1", "type": "document", "title": "FIR 112/2026", "ref_id": "doc-fir"}
    assert top["speaker"] == "Witness" and top["cross_examination_questions"]
    assert report["discarded_unverified"] == 1
    assert "212-555-0187" not in fake_contradictions["statement"] and "[PHONE]" in fake_contradictions["statement"]
    assert fake_contradictions["jurisdiction"].code == "IN"


def test_contradiction_endpoint_gathers_sources(client, fake_contradictions, audit_log, monkeypatch):
    fir_id, receipt_id, thread_id = str(uuid4()), str(uuid4()), str(uuid4())
    docs = {fir_id: ("fir.pdf", FIR), receipt_id: ("receipt.pdf", RECEIPT)}
    monkeypatch.setattr(document_service, "get_document", lambda token, did, include_text=False, org_id=None:
                        {"id": did, "filename": docs[did][0], "extracted_text": docs[did][1]} if did in docs else None)
    monkeypatch.setattr(memory_service, "load_case_memory",
                        lambda db, tid: memory_service.CaseMemory(tid, "Rao theft", "Client says the van left at 7:30 PM.", 2))

    response = client.post("/api/v1/evidence/analyze-contradictions", json={
        "statement": {"text": DEPOSITION, "label": "Deposition"},
        "document_ids": [fir_id, receipt_id], "thread_id": thread_id, "jurisdiction": "IN",
    })
    assert response.status_code == 200, response.text
    body = response.json()
    assert [s["type"] for s in body["sources"]] == ["document", "document", "case_memory"]
    assert "Rao theft" in fake_contradictions["sources"] and len(body["contradictions"]) == 2
    assert audit_log[-1]["action"] == "ANALYZE_CONTRADICTIONS"

    no_sources = client.post("/api/v1/evidence/analyze-contradictions", json={"statement": {"text": DEPOSITION}})
    assert no_sources.status_code == 422
    two_statements = client.post("/api/v1/evidence/analyze-contradictions", json={"statement": {"text": "x", "document_id": fir_id}})
    assert two_statements.status_code == 422


def test_deposition_audio_is_transcribed_and_saved(client, audit_log, monkeypatch):
    async def transcribe(audio, http_client=None):
        return {"text": "[00:00:01] Speaker 1: I left at six.", "utterances": [], "duration": 12.5}

    saved = {}
    monkeypatch.setattr(evidence, "transcribe_recording", transcribe)
    monkeypatch.setattr(document_service, "save_document", lambda token, row, org_id=None: saved.update(row) or {**row, "id": "doc-9"})
    response = client.post("/api/v1/evidence/transcribe", files={"file": ("depo.mp3", b"ID3fake-audio", "audio/mpeg")})
    assert response.status_code == 200, response.text
    assert response.json()["document_id"] == "doc-9"
    assert saved["filename"] == "depo (transcript).txt" and saved["extracted_text"].startswith("[00:00:01]")
    assert audit_log[-1]["action"] == "TRANSCRIBE_EVIDENCE"


# ============================================================================ drafting


def test_template_variables_and_slots_are_substituted():
    template = drafting_engine.get_template("legal_notice", juris.get_profile("IN"))
    rendered = drafting_engine.render(
        template,
        {"client_name": "Ravi Kumar", "recipient_name": "Suresh Patil", "compliance_days": "15", "date": "21 October 2026"},
        {"facts": "1. First fact.\n\n2. Second fact.", "demands": "3. Refund the deposit."},
    )
    text = json.dumps(rendered, ensure_ascii=False)
    assert "Under instructions from and on behalf of my client, Ravi Kumar," in text
    assert "within 15 days of receipt" in text and "{{" not in text
    facts = next(s for s in rendered["sections"] if s["heading"] == "Facts")
    assert facts["paragraphs"] == ["1. First fact.", "2. Second fact."]  # a slot becomes several paragraphs
    assert "[●Recipient's address●]" in text and "Recipient's address" in rendered["missing"]
    assert any(m.startswith("Text to draft: Numbered paragraphs on the recipient's legal liability") for m in rendered["missing"])


def test_optional_slots_can_be_omitted():
    template = drafting_engine.get_template("nda", None)
    rendered = drafting_engine.render(template, {}, {"additional_terms": "None"})
    assert rendered["sections"][-1]["heading"] == "8. Governing law"


@pytest.mark.parametrize(
    "doc_type, code, name",
    [
        ("legal_notice", "IN", "Legal notice (India)"), ("legal_notice", "US-NY", "Demand letter (US)"),
        ("legal_notice", "UK", "Letter before claim (England & Wales)"), ("bail_application", "IN", "Bail application (India, BNSS)"),
        ("bail_application", "US", "Application for bail / pretrial release"), ("affidavit", "UK", "Affidavit (England & Wales)"),
        ("nda", "IN", "Mutual non-disclosure agreement"),
    ],
)
def test_templates_are_chosen_by_jurisdiction(doc_type, code, name):
    assert drafting_engine.get_template(doc_type, juris.get_profile(code)).name == name


def test_automatic_values_follow_the_jurisdiction():
    from datetime import date

    ca = drafting_engine.automatic_variables(juris.get_profile("US-CA"), date(2026, 3, 5))
    assert ca == {"date": "March 5, 2026", "governing_law": "the laws of the State of California", "state": "California", "courts": "the State of California"}
    assert drafting_engine.automatic_variables(juris.get_profile("IN"), date(2026, 3, 5))["date"] == "5 March 2026"


def sample_draft():
    template = drafting_engine.get_template("legal_notice", juris.get_profile("IN"))
    rendered = drafting_engine.render(template, {"client_name": "Ravi Kumar", "date": "21 October 2026"},
                                      {"facts": "1. The tenancy ended on 31 March 2026.", "legal_grounds": "2. Section 70 applies.", "demands": "3. Refund Rs. 60,000."})
    return {"title": rendered["title"], "content": {"sections": rendered["sections"], "signature_block": rendered["signature_block"]}}


def test_docx_export_is_an_editable_word_document():
    data = drafting_engine.to_docx(sample_draft())
    document = Document(io.BytesIO(data))
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == "LEGAL NOTICE" and "Facts" in texts and "1. The tenancy ended on 31 March 2026." in texts
    highlighted = [r.text for p in document.paragraphs for r in p.runs if r.font.highlight_color]
    assert "[●Recipient's address●]" in highlighted and all(t.startswith("[●") for t in highlighted)
    assert document.paragraphs[0].runs[0].bold


def test_pdf_export():
    data = drafting_engine.to_pdf(sample_draft())
    text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
    assert data.startswith(b"%PDF") and "LEGAL NOTICE" in text and "Refund Rs. 60,000." in text


class FakeTable:
    """Just enough of the Supabase query builder for the drafts table."""

    def __init__(self, store):
        self.store, self.filters, self.payload, self.mode = store, {}, None, "select"

    def insert(self, row):
        self.mode, self.payload = "insert", row
        return self

    def update(self, changes):
        self.mode, self.payload = "update", changes
        return self

    def delete(self):
        self.mode = "delete"
        return self

    def select(self, *a):
        return self

    def eq(self, column, value):
        self.filters[column] = value
        return self

    def order(self, *a, **k):
        return self

    def limit(self, *a):
        return self

    def execute(self):
        rows = [r for r in self.store if all(r.get(k) == v for k, v in self.filters.items())]
        if self.mode == "insert":
            row = {**self.payload, "id": str(uuid4()), "status": "draft", "created_at": "2026-10-21T10:00:00Z", "updated_at": "2026-10-21T10:00:00Z"}
            self.store.append(row)
            return SimpleNamespace(data=[row])
        if self.mode == "update":
            for r in rows:
                r.update(self.payload)
        return SimpleNamespace(data=rows)


@pytest.fixture
def drafts_db(monkeypatch):
    store = []
    db = SimpleNamespace(table=lambda name: FakeTable(store))
    monkeypatch.setattr("app.dependencies.org.get_supabase_for_user", lambda *a, **k: db)
    return store


def test_generate_fills_the_template_from_case_context(client, drafts_db, audit_log, monkeypatch):
    seen = {}

    async def fill(spec, context, *, jurisdiction=None, instructions=None):
        seen.update(spec=spec, context=context, jurisdiction=jurisdiction)
        return DraftFill(
            variables=[NamedValue(name="recipient_name", value="Suresh Patil"), NamedValue(name="client_name", value="R. K."),
                       NamedValue(name="not_a_variable", value="x")],
            slots=[NamedValue(name="facts", value="1. My client paid a deposit; contact [PHONE]."),
                   NamedValue(name="legal_grounds", value="2. Section 70, Indian Contract Act, 1872."),
                   NamedValue(name="demands", value="3. Refund the deposit.")],
            missing_information=["Date the keys were returned"],
        )

    doc_id = str(uuid4())
    monkeypatch.setattr(anthropic_service, "is_configured", lambda: True)
    monkeypatch.setattr(anthropic_service, "fill_draft", fill)
    monkeypatch.setattr(document_service, "get_document", lambda token, did, include_text=False, org_id=None:
                        {"id": did, "filename": "lease.pdf", "extracted_text": "Clause 7: deposit of Rs. 60,000 refundable on vacating."})
    monkeypatch.setattr(memory_service, "load_case_memory",
                        lambda db, tid: memory_service.CaseMemory(tid, "Kumar", f"Client's SSN {SSN}; phone 212-555-0187.", 1))
    drafts_db.append({"thread_id": "t", "title": "x"})  # an unrelated row
    monkeypatch.setattr(
        "app.routers.drafting._case_context",
        lambda ctx, payload: (f"<case_history>Client's SSN {SSN}; phone 212-555-0187.</case_history>\n<document>deposit Rs. 60,000</document>",
                              ["Deposit of Rs. 60,000 paid on 1 March 2025"]),
    )

    response = client.post("/api/v1/drafting/generate", json={
        "doc_type": "legal_notice", "jurisdiction": "IN", "document_ids": [doc_id],
        "variables": {"client_name": "Ravi Kumar", "compliance_days": "15"},
    })
    assert response.status_code == 201, response.text
    draft = response.json()
    assert draft["template"] == "Legal notice (India)" and draft["ai_filled"] is True
    assert draft["variables"]["client_name"] == "Ravi Kumar"  # the lawyer's value beats the model's
    assert draft["variables"]["recipient_name"] == "Suresh Patil" and "not_a_variable" not in draft["variables"]
    text = json.dumps(draft["content"])
    assert "contact 212-555-0187" in text  # unmasked in the draft
    assert SSN not in seen["context"] and "212-555-0187" not in seen["context"]  # but masked for Claude
    assert "already provided by the lawyer" in seen["spec"] and seen["jurisdiction"].code == "IN"
    assert "Date the keys were returned" in draft["missing"] and "Recipient's address" in draft["missing"]
    assert audit_log[-1]["action"] == "GENERATE_DRAFT"
    basis = draft["fact_basis"]
    assert basis["provided"] == {"client_name": "Ravi Kumar", "compliance_days": "15"}  # not the model's guesses
    assert basis["case_facts"] == ["Deposit of Rs. 60,000 paid on 1 March 2025"] and basis["document_ids"] == [doc_id]
    assert "recipient_name" not in basis["variables"]

    exported = client.get(f"/api/v1/drafting/{draft['id']}/export?format=docx")
    assert exported.status_code == 200
    assert exported.headers["content-type"] == drafting_engine_docx_type()
    assert exported.headers["content-disposition"] == 'attachment; filename="legal-notice.docx"'
    assert Document(io.BytesIO(exported.content)).paragraphs[0].text == "LEGAL NOTICE"
    assert client.get(f"/api/v1/drafting/{draft['id']}/export?format=pdf").content.startswith(b"%PDF")

    content = draft["content"]
    content["sections"][0]["paragraphs"] = ["From: A. Mehta, Advocate, 12 MG Road, Pune"]
    patched = client.patch(f"/api/v1/drafting/{draft['id']}", json={"content": content})
    assert "Advocate's office address" not in patched.json()["missing"]


def drafting_engine_docx_type() -> str:
    from app.routers.drafting import DOCX_MEDIA_TYPE

    return DOCX_MEDIA_TYPE


def test_clause_revision_returns_an_inline_diff(client, monkeypatch):
    async def revise(text, instruction, *, jurisdiction=None, context=None):
        assert "legal@acme.com" not in text
        return ClauseRevision(revised_text=text.replace("30 days", "45 days"), explanation="Extended the payment period.")

    monkeypatch.setattr(anthropic_service, "is_configured", lambda: True)
    monkeypatch.setattr(anthropic_service, "revise_clause", revise)
    response = client.post("/api/v1/drafting/revise-clause", json={
        "text": "2. Payment within 30 days; notices to legal@acme.com.", "instruction": "Give the client more time to pay"})
    body = response.json()
    assert body["revised_text"] == "2. Payment within 45 days; notices to legal@acme.com."
    assert {"op": "delete", "text": "30"} in body["inline_diff"] and {"op": "insert", "text": "45"} in body["inline_diff"]


def test_unknown_document_types_are_rejected(client):
    assert client.post("/api/v1/drafting/generate", json={"doc_type": "will"}).status_code == 422


def test_drafts_are_self_checked_marked_unreviewed_and_approved(client, drafts_db, audit_log, monkeypatch):
    monkeypatch.setattr(document_service, "get_document", lambda token, did, include_text=False, org_id=None:
                        {"id": did, "filename": "cheque.pdf", "extracted_text": "Cheque No. 004512 for Rs. 2,50,000 dated 05/01/2026"})
    draft_id = str(uuid4())
    drafts_db.append({
        "id": draft_id, "title": "Legal notice", "status": "draft", "approved_at": None, "approved_by": None,
        "content": {"sections": [{"heading": "Facts", "paragraphs": [
            "Your cheque No. 004512 dated 05/01/2026 for Rs. 2,05,000 was dishonoured on 12 January 2026.",
            "You are liable under Section 138 of the Negotiable Instruments Act, 1881 and Section 999 BNS.",
        ]}], "signature_block": ["For Mehta & Co."]},
        "fact_basis": {"variables": {"client_name": "Ravi Kumar"}, "provided": {"client_name": "Ravi Kumar"},
                       "notes": "Returned unpaid on 12 January 2026", "case_facts": [], "document_ids": ["d1"]},
    })

    report = client.get(f"/api/v1/drafting/{draft_id}/checklist").json()
    mismatches = {(m["kind"], m["value"]) for m in report["mismatches"]}
    assert ("amount", "205000") in mismatches  # the document says 2,50,000: digits swapped
    assert ("provided_value", "Ravi Kumar") in mismatches  # the client's name is missing from the draft
    assert not any(kind == "date" for kind, _ in mismatches) and not any(kind == "reference" for kind, _ in mismatches)
    assert [c["text"] for c in report["unresolved_citations"]] == ["Section 999 BNS"]
    assert report["approved"] is False and report["clear"] is False

    # Unapproved exports are marked on every page; unresolved citations can be swapped for a placeholder.
    exported = client.get(f"/api/v1/drafting/{draft_id}/export?format=docx&unresolved=remove")
    document = Document(io.BytesIO(exported.content))
    assert document.sections[0].header.paragraphs[0].text.startswith("AI DRAFT · UNREVIEWED")
    body = "\n".join(p.text for p in document.paragraphs)
    assert "Section 999 BNS" not in body and "[●citation to be confirmed●]" in body and "Section 138" in body
    assert audit_log[-1]["metadata"]["unresolved_citations"] == ["Section 999 BNS"]
    assert audit_log[-1]["metadata"]["unresolved_choice"] == "remove" and audit_log[-1]["metadata"]["approved"] is False

    approved = client.post(f"/api/v1/drafting/{draft_id}/approve")
    assert approved.status_code == 200 and approved.json()["approved_at"]
    assert audit_log[-1]["action"] == "APPROVE_DRAFT"
    exported = client.get(f"/api/v1/drafting/{draft_id}/export?format=docx")
    assert Document(io.BytesIO(exported.content)).sections[0].header.paragraphs[0].text == ""
    assert client.get(f"/api/v1/drafting/{draft_id}/export?format=pdf").content.startswith(b"%PDF")
    assert client.delete(f"/api/v1/drafting/{draft_id}/approve").json()["approved_at"] is None
